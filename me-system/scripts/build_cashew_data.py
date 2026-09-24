"""Build src/data/cashew.json from the Cashew Policy M&E source files.

Usage:
    python3 scripts/build_cashew_data.py <folder with the source files>

Source files (matched by name fragment, not committed to the repository):
    Cashew_Dashboard__Database_Action_Level_v5_*.xlsx   action-level workbook
    Outcome_Dashboard_v12.xlsx                          outcome-level workbook
    CASHEW_INDICATOR_REPORT_Kobo_XLSForm_v10_*.xlsx     Kobo XLSForm
    *Cashew_MTR_Report_Final_Draft*.docx                mid-term review report

Requires: openpyxl, python-docx.
The output holds aggregates, indicator definitions, form questions and
report text only. No respondent names or phone numbers are exported.
"""

import glob
import json
import os
import re
import sys
from collections import OrderedDict, defaultdict

import docx
import openpyxl


def find(folder, fragment):
    hits = [p for p in glob.glob(os.path.join(folder, "*")) if fragment in os.path.basename(p)]
    if not hits:
        sys.exit(f"Missing source file containing '{fragment}' in {folder}")
    return hits[0]


def sheet_rows(wb, name):
    rows = wb[name].iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    for r in rows:
        if r and any(v not in (None, "") for v in r):
            yield dict(zip(header, r))


def clean(text):
    if text is None:
        return ""
    return re.sub(r"[ \t]+", " ", str(text).replace("​", "").replace("‐", "-")).strip()


def num(v):
    if v in (None, ""):
        return None
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return int(f) if f.is_integer() else round(f, 4)


# Short display codes used across the MoC dashboards.
SHORT = {
    "maff": "MAFF", "moe": "MoE", "moc": "MoC", "ovop": "OVOP", "moi": "MoI", "mef": "MEF",
    "nbc": "NBC", "mowram": "MOWRAM", "cmaa": "CMAA", "mrd": "MRD", "cdc": "CDC",
    "misti": "MISTI", "mot": "MoT", "gdce": "GDCE", "mome": "MoME", "mptc": "MPTC", "mpwt": "MPWT",
}

# Action clusters as grouped in the MTR report (tables 4, 5 and 6).
def cluster_of(action_no):
    if action_no <= 17:
        return "production"
    if action_no <= 28:
        return "processing"
    return "export"


# Workbook v5 DASHBOARD_DATA pillar mapping, kept for the data note.
def workbook_pillar(action_no):
    if action_no <= 16:
        return "production"
    if action_no <= 32:
        return "processing"
    return "export"


def method_of(answer_type, calculation):
    if answer_type.startswith("select_one"):
        return "milestone", None
    m = re.search(r"\((\d+(?:\.\d+)?) div \$\{", calculation)
    if m:
        return "inverse_time", num(m.group(1))
    m = re.search(r"div (\d+(?:\.\d+)?)\)", calculation)
    if m:
        return "count_to_target", num(m.group(1))
    return "percent_complete", 100


def main(folder):
    out = OrderedDict()

    # ---------- Action-level workbook ----------
    wb = openpyxl.load_workbook(find(folder, "Database_Action_Level"), read_only=True, data_only=True)

    master = list(sheet_rows(wb, "02_INDICATOR_MASTER"))
    latest = {r["Indicator_ID"]: r for r in sheet_rows(wb, "05_LATEST_VALID_DATA") if r.get("Reporting_Year") == 2025}
    thresholds = [
        {"year": int(r["Reporting_Year"]), "largely": num(r["Largely_Achieved_Min_%"]), "fully": num(r["Fully_Achieved_Min_%"])}
        for r in sheet_rows(wb, "10_STATUS_THRESHOLDS")
    ]
    action_prog = {
        int(r["Action_No"]): r for r in sheet_rows(wb, "07_ACTION_PROGRESS") if r.get("Reporting_Year") == 2025
    }
    ministry_summary = {
        r["Ministry_Code"]: r for r in sheet_rows(wb, "06_MINISTRY_SUMMARY") if r.get("Reporting_Year") == 2025
    }

    # ---------- Kobo XLSForm (Khmer labels, constraints) ----------
    xf = openpyxl.load_workbook(find(folder, "Kobo_XLSForm"), read_only=True, data_only=True)
    survey = list(sheet_rows(xf, "survey"))
    by_name = {r["name"]: r for r in survey if r.get("name")}
    settings = next(sheet_rows(xf, "settings"))
    choices = [
        {"list": r["list_name"], "name": str(r["name"]), "en": clean(r["label::English (en)"]), "km": clean(r["label::Khmer (km)"])}
        for r in sheet_rows(xf, "choices")
    ]

    ministries = OrderedDict()
    for c in choices:
        if c["list"] == "ministry" and c["name"] != "admin":
            m = re.match(r"^[^(]+\((.*)\)$", c["en"])
            mk = re.match(r"^[^(]+\((.*)\)$", c["km"])
            ministries[c["name"]] = {
                "code": c["name"],
                "short": SHORT.get(c["name"], c["name"].upper()),
                "name": m.group(1) if m else c["en"],
                "nameKm": mk.group(1) if mk else c["km"],
            }

    indicators = []
    for r in master:
        iid = r["Indicator_ID"]
        n = int(r["Indicator_No"])
        key = f"ind_{n:03d}"
        method, target = method_of(r["Answer_Type"], r["Calculation"])
        value_row = by_name.get(f"{key}_value", {})
        header_row = by_name.get(f"{key}_header", {})
        lv = latest.get(iid, {})
        indicators.append(
            {
                "id": iid,
                "no": n,
                "code": f"AI-{n:03d}",
                "action": int(r["Action_No"]),
                "ministry": r["Assigned_Ministry_Code"],
                "label": clean(r["Indicator_Label_EN"]),
                "labelKm": clean(str(header_row.get("label::Khmer (km)", "")).split("\n")[0].split("៖", 1)[-1]),
                "targetText": clean(r["Target_Text"]) or "100% by 2027",
                "target": target,
                "method": method,
                "answerType": r["Answer_Type"],
                "calculation": clean(r["Calculation"]),
                "question": clean(r["Question_EN"]),
                "questionKm": clean(value_row.get("label::Khmer (km)")),
                "constraint": clean(value_row.get("constraint")),
                "constraintMsg": clean(value_row.get("constraint_message::English (en)")),
                "kobo": {
                    "group": r["Group_Name"],
                    "value": r["Value_Field"],
                    "percentage": r["Percentage_Field"],
                    "followup": r["Followup_Field"],
                    "evidence": r["Evidence_Field"],
                    "feedback": r["Feedback_Field"],
                },
                "y2025": {
                    "value": lv.get("Reported_Value") if isinstance(lv.get("Reported_Value"), str) else num(lv.get("Reported_Value")),
                    "actualPct": num(lv.get("KPI_Actual_%")),
                    "cappedPct": num(lv.get("KPI_Capped_%")),
                    "status": lv.get("Progress_Status") or None,
                },
            }
        )

    # ---------- MTR report ----------
    doc = docx.Document(find(folder, "MTR_Report"))

    def table_rows(t):
        rows = []
        for row in t.rows:
            cells = []
            for c in row.cells:
                x = clean(c.text.replace("\n", " / "))
                if not cells or cells[-1] != x:
                    cells.append(x)
            rows.append(cells)
        return rows

    mtr_actions = {}
    for t in doc.tables[4:7]:
        for cells in table_rows(t)[1:]:
            cells = [c for c in cells if c.strip(" /") != ""]
            if len(cells) < 4:
                continue
            activity, progress, summary, responsible = cells[0], cells[1], cells[2], cells[3]
            nums = [int(x) for x in re.findall(r"(?:^|/ )(\d{1,2})\.\s", activity)]
            text = re.sub(r"^\d{1,2}\.\s*", "", activity)
            parts = re.split(r" / (?=\d{1,2}\.\s)", text)
            for i, a in enumerate(nums or []):
                mtr_actions[a] = {
                    "title": re.sub(r"^\d{1,2}\.\s*", "", parts[i] if i < len(parts) else parts[-1]).strip(" /"),
                    "mtrProgress": re.sub(r"(\s*/\s*)+$", "", progress),
                    "mtrSummary": summary,
                    "responsible": responsible.replace(" / ", ", ").strip(" /,"),
                }

    paragraphs = [(p.style.name, clean(p.text)) for p in doc.paragraphs if clean(p.text)]

    def section(start, stop):
        take, buf = False, []
        for style, text in paragraphs:
            if style.startswith("Heading") and text.startswith(start):
                take = True
                continue
            if take and style.startswith("Heading") and stop and text.startswith(stop):
                break
            if take and not style.startswith("Heading"):
                buf.append(text)
        return buf

    recs = OrderedDict()
    for title, start, stop in [
        ("Production", "5.1", "5.2"),
        ("Investment in processing", "5.2", "5.3"),
        ("Working capital for processors", "5.3", "5.4"),
        ("Price competitiveness of processors", "5.4", "5.5"),
        ("Exports to key markets", "5.5", "5.6"),
        ("Coordination and M&E", "5.6", "6."),
    ]:
        recs[title] = section(start, stop)

    exec_summary = section("EXECUTIVE SUMMARY", "Priority Activities")
    out["mtr"] = {
        "title": "Mid-Term Review of the National Cashew Policy 2022–2027",
        "status": "Final draft",
        "date": "2026-05-25",
        "asOf": "mid-2025",
        "goalProgress": [p for p in exec_summary if p[:2] in ("A.", "B.", "C.")],
        "criteria": {
            k: next((exec_summary[i + 1] for i, p in enumerate(exec_summary) if p == k), "")
            for k in ["Relevance", "Coherence", "Effectiveness", "Efficiency", "Impact", "Sustainability"]
        },
        "priorities": section("Priority Activities", "1. INTRODUCTION"),
        "recommendations": recs,
        "nextSteps": section("6. CONCLUSION", "List of Abbreviations")[-4:],
        "clusterCounts": {
            "production": {"actions": 17, "fully": 6, "largely": 8, "limited": 3},
            "processing": {"actions": 11, "fully": 3, "largely": 6, "limited": 2},
            "export": {"actions": 16, "fully": 6, "largely": 9, "limited": 1},
            "all": {"actions": 44, "fully": 15, "largely": 23, "limited": 6},
        },
        "method": "Document review, 22 key informant interviews, visits to six processing factories in Kampong Thom and Kampong Cham, validation workshop. MTR rule: < 50% limited, ≥ 50% largely, completed = fully achieved.",
    }

    # ---------- Actions ----------
    actions = []
    for a in range(1, 45):
        inds = [i for i in indicators if i["action"] == a]
        owners = list(OrderedDict.fromkeys(i["ministry"] for i in inds))
        ap = action_prog.get(a, {})
        mt = mtr_actions.get(a, {})
        actions.append(
            {
                "no": a,
                "code": f"ACT-{a:02d}",
                "title": mt.get("title") or inds[0]["label"],
                "cluster": cluster_of(a),
                "workbookPillar": workbook_pillar(a),
                "lead": owners[0],
                "ministries": owners,
                "responsible": mt.get("responsible", ""),
                "indicators": [i["id"] for i in inds],
                "y2025": {
                    "avgActualPct": num(ap.get("Avg_Actual_%")),
                    "avgCappedPct": num(ap.get("Avg_Capped_%")),
                    "fully": num(ap.get("Fully_Achieved")) or 0,
                    "largely": num(ap.get("Largely_Achieved")) or 0,
                    "limited": num(ap.get("Limited_Progress")) or 0,
                },
                "mtrProgress": mt.get("mtrProgress", ""),
                "mtrSummary": mt.get("mtrSummary", ""),
            }
        )

    for code, m in ministries.items():
        s = ministry_summary.get(code, {})
        m["y2025"] = {
            "assigned": num(s.get("Assigned_Indicators")),
            "reported": num(s.get("Reported_Indicators")),
            "avgActualPct": num(s.get("Avg_Actual_%")),
            "avgCappedPct": num(s.get("Avg_Capped_%")),
            "fully": num(s.get("Fully_Achieved")) or 0,
            "largely": num(s.get("Largely_Achieved")) or 0,
            "limited": num(s.get("Limited_Progress")) or 0,
        }

    # ---------- Kobo form structure (sections per ministry) ----------
    sections = []
    current = None
    for r in survey:
        t, name = str(r.get("type") or ""), r.get("name") or ""
        if t == "begin_group" and name.startswith("grp_"):
            current = {
                "name": name,
                "ministry": name[4:],
                "label": clean(r.get("label::English (en)")),
                "labelKm": clean(r.get("label::Khmer (km)")),
                "relevant": clean(r.get("relevant")),
                "indicators": [],
            }
            sections.append(current)
        elif t == "begin_group" and name.startswith("g_ind_") and current is not None:
            current["indicators"].append("Indicator_" + name[6:])
    shared = [by_name[k] for k in ["ind_001_followup", "ind_001_evidence", "ind_001_feedback"]]
    out["koboForm"] = {
        "title": clean(settings.get("form_title")),
        "formId": clean(settings.get("form_id")),
        "version": str(settings.get("version")),
        "defaultLanguage": clean(settings.get("default_language")),
        "rows": len(survey),
        "respondentFields": [
            {"name": by_name[k]["name"], "type": by_name[k]["type"], "en": clean(by_name[k]["label::English (en)"]), "km": clean(by_name[k]["label::Khmer (km)"]), "hint": clean(by_name[k].get("hint::English (en)"))}
            for k in ["respondent_name", "phone_number", "reporting_year", "reporting_ministry"]
        ],
        "perIndicatorFields": [
            {"suffix": "_value", "type": "integer / decimal / select_one", "en": "Indicator question (reported value)", "km": "", "hint": "Hint shows the 2025 reported progress."},
            {"suffix": "_percentage", "type": "calculate", "en": "Percentage of the 2027 target (auto-calculated)", "km": "", "hint": ""},
        ]
        + [
            {"suffix": "_" + s["name"].split("_", 2)[2], "type": s["type"], "en": clean(s["label::English (en)"]), "km": clean(s["label::Khmer (km)"]), "hint": clean(s.get("hint::English (en)")), "required": s.get("required") == "yes"}
            for s in shared
        ],
        "sections": sections,
        "choices": choices,
    }

    # ---------- Outcome workbook ----------
    ow = openpyxl.load_workbook(find(folder, "Outcome_Dashboard_v12"), read_only=True, data_only=True)
    dash = list(ow["02_DASH_DATA"].iter_rows(values_only=True))
    master_o = {r["Code"]: r for r in sheet_rows_offset(ow, "03_INDICATORS_MASTER", "Code")}
    years = [2022, 2023, 2024, 2025, 2026, 2027]
    outcome = []
    for row in dash[4:17]:
        code = row[1]
        mo = master_o.get(code, {})
        outcome.append(
            {
                "code": code,
                "title": clean(row[2]),
                "unit": clean(row[3]),
                "direction": "increase" if row[4] == "Higher" else "decrease",
                "series": {str(y): num(row[5 + i]) for i, y in enumerate(years)},
                "source": clean(mo.get("Primary reporter")),
                "tool": clean(mo.get("Data tool")),
                "area": clean(row[12]),
                "frequency": clean(mo.get("Frequency")),
                "disaggregation": clean(mo.get("Disaggregation")),
                "definition": clean(mo.get("Definition / formula")),
                "note": clean(row[13]),
                "testData": clean(mo.get("Data tool")) == "Kobo",
            }
        )
    out["outcomeIndicators"] = outcome

    markets = defaultdict(lambda: defaultdict(lambda: {"kernel_t": 0.0, "kernel_usd": 0.0, "processed_t": 0.0, "processed_usd": 0.0, "rcn_t": 0.0, "rcn_eq_t": 0.0}))
    for r in sheet_rows(ow, "05_MARKETS"):
        y, p, d = r.get("Year"), r.get("Product type"), r.get("Destination")
        if not isinstance(y, (int, float)) or not d:
            continue
        m = markets[int(y)][clean(d)]
        q, v, eq = num(r.get("Export quantity (t)")) or 0, num(r.get("Export value (USD)")) or 0, num(r.get("RCN equivalent (t)")) or 0
        key = {"RCN": "rcn", "Kernel": "kernel", "Processed": "processed"}.get(p)
        if key == "rcn":
            m["rcn_t"] += q
        elif key:
            m[f"{key}_t"] += q
            m[f"{key}_usd"] += v
        m["rcn_eq_t"] += eq
    top = {}
    for y, dests in markets.items():
        rows = sorted(
            ({"destination": d, **{k: round(v, 2) for k, v in vals.items()}} for d, vals in dests.items() if vals["kernel_t"] + vals["processed_t"] > 0),
            key=lambda x: -(x["kernel_t"] * 4 + x["processed_t"] * 4.1),
        )
        top[str(y)] = rows[:10]
    out["topMarkets"] = top

    out["production"] = [
        {"year": int(r["Year"]), "area_ha": num(r["Productive area (ha)"]), "production_t": num(r["Annual production (t RCN)"]), "source": clean(r["Data source"])}
        for r in sheet_rows(ow, "04_PRODUCTION")
        if num(r.get("Productive area (ha)"))
    ]
    out["income"] = [
        {"year": int(r["Year"]), "price_khr": num(r["Farmgate price (KHR/kg)"]), "fx": num(r["Exchange rate (KHR/USD)"]), "income_usd": num(r["Est. farmer income (USD) [S1]"])}
        for r in sheet_rows(ow, "06_INCOME")
        if num(r.get("Farmgate price (KHR/kg)"))
    ]

    out["ministries"] = list(ministries.values())
    out["actions"] = actions
    out["actionIndicators"] = indicators
    out["thresholds"] = thresholds

    dest = os.path.join(os.path.dirname(__file__), "..", "src", "data", "cashew.json")
    os.makedirs(os.path.dirname(dest), exist_ok=True)
    with open(dest, "w", encoding="utf-8") as fh:
        json.dump(out, fh, ensure_ascii=False, indent=1)
    print(f"Wrote {dest}: {len(ministries)} ministries, {len(actions)} actions, {len(indicators)} action indicators, {len(outcome)} outcome indicators, MTR action rows {len(mtr_actions)}")


def sheet_rows_offset(wb, name, first_header):
    """Rows below the first row whose first cell equals first_header."""
    rows = wb[name].iter_rows(values_only=True)
    for r in rows:
        if r and r[0] == first_header:
            header = [str(h).strip() if h is not None else "" for h in r]
            break
    for r in rows:
        if r and any(v not in (None, "") for v in r):
            yield dict(zip(header, r))


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
